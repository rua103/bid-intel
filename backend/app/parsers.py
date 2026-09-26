from __future__ import annotations

import importlib
import importlib.util
import io
import re
import shutil
import struct
import subprocess
import tempfile
import unicodedata
import zipfile
import zlib
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import settings
from app.file_formats import inspect_content
from app.legacy_documents import DocumentConversionError, convert_doc, libreoffice_executable
from app.schemas import ItemCandidate, NoticeMetadata, ParticipantCandidate

MAX_ARCHIVE_DEPTH = 3
MAX_ARCHIVE_FILES = 2_000
MAX_EXPANDED_BYTES = 200 * 1024 * 1024

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "package_code": ("采购包编号", "采购包号", "包编号", "包号", "标包编号", "标段编号", "标段号"),
    "product_name": (
        "标的名称",
        "标的物名称",
        "采购标的",
        "采购标的名称",
        "标项名称",
        "采购内容",
        "产品服务名称",
        "产品名称",
        "货物名称",
        "服务名称",
        "货物服务名称",
        "名称",
        "品名",
        "标的物",
    ),
    "category": ("品目名称", "品目", "采购品目", "品目分类", "类别"),
    "brand": ("品牌", "品牌如有", "产品供应商", "品牌产品供应商"),
    "model": ("规格型号", "规格型号如有", "规格说明", "规格", "型号"),
    "quantity": ("数量", "数量单位", "采购数量", "采购数量单位"),
    "unit_price": ("单价", "单价元"),
    "total_price": (
        "总价", "合价", "总金额", "成交金额", "中标金额", "中标成交金额", "金额",
    ),
}

# These aliases deliberately describe an organization that responded to a
# procurement.  Generic columns such as "单位" and "名称" are not accepted:
# they also occur in item, buyer, and agent tables.
PARTICIPANT_ORG_ALIASES = (
    "投标人名称", "投标单位名称", "供应商名称", "供应商信息", "响应供应商",
    "供应商全称", "响应人名称", "响应单位名称", "投标人", "投标单位",
    "供应商", "响应人", "响应单位", "候选供应商",
)
PARTICIPANT_PACKAGE_ALIASES = (
    "采购包编号", "采购包号", "合同包号", "包号", "标包号", "标包", "标段号", "合同包", "采购包",
)
PARTICIPANT_OUTCOME_ALIASES = (
    "是否中标", "是否成交", "中标状态", "成交状态", "中标情况", "成交情况",
    "中标结果", "成交结果",
)
PARTICIPANT_AWARD_AMOUNT_ALIASES = (
    "中标成交金额", "中标金额", "成交金额", "中标总金额", "成交总金额",
)
PARTICIPANT_UNSUCCESSFUL_ALIASES = (
    "未中标原因", "未中标成交原因", "未成交原因", "未成交中标原因", "落标原因",
)
PARTICIPANT_QUALIFICATION_ALIASES = ("资格审查", "资格审查结果", "资格性审查", "资格性审查结果")
PARTICIPANT_COMPLIANCE_ALIASES = ("符合性审查", "符合性审查结果", "符合性检查")
PARTICIPANT_REVIEW_SIGNALS = (
    "资格性审查", "资格审查", "符合性审查", "符合性检查", "评审总得分",
    "综合得分", "评审得分", "得分排名", "推荐排名", "投标报价", "响应报价",
    "评标结果", "未中标原因", "未中标成交原因", "未成交原因", "落标原因",
)
PARTICIPANT_CONTEXT_SIGNALS = (
    "投标人名单", "投标单位名单", "投标情况", "参与投标", "开标一览",
    "报价一览", "资格性审查", "资格审查", "符合性审查", "评审得分",
    "候选供应商", "中标候选人", "成交候选人", "评标结果", "未中标供应商",
    "未成交供应商",
)


@dataclass(frozen=True)
class SourceDocument:
    filename: str
    content: bytes


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value) if value is not None else "").strip()


def _header_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", _clean(value)).lower()
    text = re.sub(r"[\s（）()【】\[\]：:，,、/\\_-]", "", text)
    text = text.replace("人民币", "")
    text = re.sub(r"(?:亿元|万元|元)$", "", text)
    return text


def _column_map(header: Iterable[object]) -> dict[str, int]:
    cells = [_header_key(cell) for cell in header]
    result: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        normalized_aliases = [_header_key(alias) for alias in aliases]
        # Match whole labels, never prose containing a header word. Alias order
        # gives 标的名称 priority over the broader 标项名称 when both are present.
        for alias in normalized_aliases:
            for index, cell in enumerate(cells):
                if cell and cell == alias:
                    result[field] = index
                    break
            if field in result:
                break
    return result


def _requirement_header(header: list[str], columns: dict[str, int]) -> bool:
    labels = {_header_key(cell) for cell in header}
    return bool(labels & {'简要技术要求', '采购需求', '技术要求', '技术参数要求'}
                and not {'brand', 'model', 'unit_price', 'total_price'} & columns.keys())


def _decimal(value: str) -> Decimal | None:
    if not value:
        return None
    text = value.replace(",", "").replace("，", "").replace("￥", "").replace("¥", "").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return Decimal(match.group(0))
    except InvalidOperation:
        return None


def _quantity(value: str) -> tuple[Decimal | None, str | None]:
    if not value:
        return None, None
    match = re.search(r"(-?\d+(?:[,.，]\d+)*)\s*(.*)", value)
    if not match:
        return _decimal(value), None
    return _decimal(match.group(1)), match.group(2).strip() or None


def _money(value: str, header: str = "") -> Decimal | None:
    """Accept a single RMB amount; a percentage or reference is not a price."""
    text = unicodedata.normalize("NFKC", value)
    text = re.sub(r"[\s,，￥¥()]|人民币", "", text)
    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)(亿元|万元|元)?", text)
    if not match:
        return None
    unit = match.group(2) or ("亿元" if "亿元" in header else
                            "万元" if "万元" in header else "元")
    return Decimal(match.group(1)) * {"元": 1, "万元": 10_000, "亿元": 100_000_000}[unit]


def _item_header(columns: dict[str, int]) -> bool:
    return (len(set(columns.values())) >= 3
            and ("product_name" in columns or "category" in columns))


def parse_item_tables(
    rows: list[list[str]], *, source_file: str, table_index: int
) -> list[ItemCandidate]:
    """Map a tabular procurement-item section to the seven contest fields."""
    candidates: list[ItemCandidate] = []
    header_index: int | None = None
    columns: dict[str, int] = {}
    for index, row in enumerate(rows[:10]):
        current = _column_map(row)
        # A flattened heading in one merged cell can contain many field words.
        # It is not a usable table header unless those fields occupy separate columns.
        if _item_header(current):
            header_index, columns = index, current
            break
    if header_index is None:
        return candidates
    header = rows[header_index]
    requirement = _requirement_header(header, columns)
    quotation = any('供应商报价成交明细' in _clean(''.join(row))
                    for row in rows[:header_index])
    if any("供应商" in cell for cell in header) and not (
        {"brand", "model", "quantity", "unit_price"} & columns.keys()
    ):
        # 包号/采购内容/供应商/中标金额 describes a package award, not
        # item quantities and prices. Do not add its total beside detailed items.
        return candidates

    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = [_clean(cell) for cell in row]
        if not any(values):
            continue
        repeated = _column_map(values)
        if _item_header(repeated):
            columns, header = repeated, values
            requirement = _requirement_header(header, columns)
            continue
        if requirement:
            continue

        def get_value(
            field: str, row: list[str] = values, column_map: dict[str, int] = columns
        ) -> str:
            col = column_map.get(field)
            return row[col] if col is not None and col < len(row) else ""

        product_name = get_value("product_name")
        category = get_value("category")
        quantity, quantity_unit = _quantity(get_value("quantity"))
        extra = {}
        if "package_code" in ItemCandidate.model_fields:
            extra["package_code"] = get_value("package_code") or "default"
            # Some official quotation sheets put 包1 in the 序号 column.
            # Bare numeric row indices must never become package identifiers.
            if extra['package_code'] == 'default':
                for col, label in enumerate(header):
                    if _header_key(label) == '序号' and col < len(values) and re.fullmatch(
                        r'(?:合同包|采购包|包)\s*\d+', values[col]
                    ):
                        extra['package_code'] = values[col]
        total = _money(get_value('total_price'),
                       header[columns['total_price']] if 'total_price' in columns else '')
        # 报价 alone is ambiguous; accept it only in an explicit 成交明细 sheet
        # with quantity/unit-price columns and a numeric row-level amount.
        if total is None and quotation and {'quantity', 'unit_price'} <= columns.keys():
            for col, label in enumerate(header):
                if _header_key(label) == '报价' and col < len(values):
                    total = _money(values[col], label)
        candidate = ItemCandidate(
            product_name=product_name or None,
            category=category or None,
            brand=get_value("brand") or None,
            model=get_value("model") or None,
            quantity=quantity,
            quantity_unit=quantity_unit,
            unit_price=_money(get_value("unit_price"),
                              header[columns["unit_price"]] if "unit_price" in columns else ""),
            total_price=total,
            source_file=source_file,
            source_location=f"table:{table_index}/row:{row_index}",
            source_evidence=" | ".join(value for value in values if value),
            **extra,
        )
        if _is_item_data_row(values, columns, candidate):
            candidates.append(candidate)
    return candidates


def _participant_column_map(header: Iterable[object]) -> dict[str, int]:
    cells = [_header_key(cell) for cell in header]
    aliases: dict[str, tuple[str, ...]] = {
        "organization": PARTICIPANT_ORG_ALIASES,
        "package": PARTICIPANT_PACKAGE_ALIASES,
        "outcome": PARTICIPANT_OUTCOME_ALIASES,
        "award_amount": PARTICIPANT_AWARD_AMOUNT_ALIASES,
        "unsuccessful_reason": PARTICIPANT_UNSUCCESSFUL_ALIASES,
        "qualification": PARTICIPANT_QUALIFICATION_ALIASES,
        "compliance": PARTICIPANT_COMPLIANCE_ALIASES,
    }
    result: dict[str, int] = {}
    for field, field_aliases in aliases.items():
        normalized_aliases = [_header_key(alias) for alias in field_aliases]
        for alias in normalized_aliases:
            for index, cell in enumerate(cells):
                if cell and cell == alias:
                    result[field] = index
                    break
            if field in result:
                break
    return result


def _participant_package_code(value: str | None) -> str:
    value = _clean(value or "")
    if not value:
        return "default"
    compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", value))
    named_match = re.search(r"包名[:：]?([A-Za-z0-9一二三四五六七八九十]+)", compact)
    if named_match:
        return f"包{_normalize_package_number(named_match.group(1))}"
    match = re.search(r"(?:合同|采购|标段|标包)?包(?:编号|号)?([A-Za-z0-9一二三四五六七八九十]+)", compact)
    if match:
        return f"包{_normalize_package_number(match.group(1))}"
    if re.fullmatch(r"\d+", compact):
        return f"包{compact}"
    return compact


def _normalize_package_number(value: str) -> str:
    chinese_numbers = {
        "一": "1", "二": "2", "三": "3", "四": "4", "五": "5",
        "六": "6", "七": "7", "八": "8", "九": "9", "十": "10",
    }
    return chinese_numbers.get(value, value)


def _explicit_package_hint(rows: list[list[str]]) -> str | None:
    """Read a package label only from this table's own preamble."""
    for row in rows[:10]:
        joined = _clean(" ".join(row))
        named_match = re.search(
            r"包名\s*[:：]?\s*([A-Za-z0-9一二三四五六七八九十]+)",
            unicodedata.normalize("NFKC", joined),
        )
        if named_match:
            return f"包{named_match.group(1)}"
        match = re.search(
            r"(?:合同|采购|标段|标包)?包\s*(?:编号|号)?\s*([A-Za-z0-9一二三四五六七八九十]+)",
            unicodedata.normalize("NFKC", joined),
        )
        if match:
            return f"包{match.group(1)}"
    return None


def _participant_outcome(
    row: list[str], columns: dict[str, int], *, award_amount: Decimal | None,
) -> str:
    """Use explicit result text only; rank and score values never determine outcome."""
    if "unsuccessful_reason" in columns:
        column = columns["unsuccessful_reason"]
        reason = _header_key(row[column] if column < len(row) else "")
        if reason and reason not in {"无", "暂无", "不适用", "-", "—", "/"}:
            return "nonwinner"
    if "outcome" in columns:
        column = columns["outcome"]
        status = _header_key(row[column] if column < len(row) else "")
        if any(token in status for token in ("未中标", "非中标", "未成交", "未获中标", "落标")):
            return "nonwinner"
        if status in {"否", "未是"}:
            return "nonwinner"
        if any(token in status for token in ("中标", "成交")) and "候选" not in status:
            return "winner"
        if status in {"是", "已中标", "已成交"}:
            return "winner"
    for field in ("qualification", "compliance"):
        if field in columns:
            column = columns[field]
            status = _header_key(row[column] if column < len(row) else "")
            if any(token in status for token in ("不通过", "未通过", "不合格", "不符合")):
                return "nonwinner"
    if award_amount is not None:
        return "winner"
    return "unknown"


def _direct_winner_key_values(
    rows: list[list[str]], *, source_file: str, table_index: int, context: str = "",
) -> list[ParticipantCandidate]:
    """Read explicit winner fields in CCGP two-column/key-value result tables."""
    winner_labels = {
        _header_key(label) for label in
        ("中标供应商", "成交供应商", "中标人", "成交人", "中标单位", "成交单位")
    }
    amount_labels = {_header_key(label) for label in PARTICIPANT_AWARD_AMOUNT_ALIASES}
    package_hint = _explicit_package_hint([[context], *rows])
    found: list[ParticipantCandidate] = []
    for row_index, raw_row in enumerate(rows, start=1):
        values = [_clean(value) for value in raw_row]
        winner_values = [
            values[index + 1]
            for index, value in enumerate(values[:-1])
            if _header_key(value) in winner_labels
        ]
        if not winner_values:
            continue
        organization = next((value for value in winner_values if value), "")
        if not organization or _invalid_participant_name(organization):
            continue
        amount: Decimal | None = None
        for index, value in enumerate(values[:-1]):
            if _header_key(value) in amount_labels:
                amount_header = value
                amount = _money(values[index + 1], amount_header)
                if amount is not None:
                    break
        found.append(ParticipantCandidate(
            organization_name=organization,
            package_code=_participant_package_code(package_hint),
            outcome="winner",
            award_amount=amount,
            source_file=source_file,
            source_location=f"table:{table_index}/row:{row_index}",
            source_evidence=" | ".join(value for value in values if value),
            extraction_method="participant_table_header_mapping",
            confidence=0.8,
        ))
    return found


def _invalid_participant_name(value: str) -> bool:
    normalized = _header_key(value)
    invalid_labels = {
        "无", "暂无", "无供应商", "无投标人", "名称", "供应商名称", "投标人名称",
        "投标单位名称", "供应商", "投标人", "响应供应商", "联系方式", "联系人",
        "联系电话", "联系地址", "企业类型", "采购单位", "采购人", "采购代理机构",
        "代理机构", "中标金额", "成交金额", "未中标原因", "备注", "说明",
    }
    non_org_phrases = (
        "综合得分", "评审得分", "未中标", "未成交", "不通过", "未通过", "不合格",
        "不符合", "得分较低", "报价低于", "报价较低", "报价偏高", "响应报价",
        "审查结果", "未满足",
    )
    return (
        not value or normalized in invalid_labels
        or any(_header_key(phrase) in normalized for phrase in non_org_phrases)
    )


def parse_participant_tables(
    rows: list[list[str]], *, source_file: str, table_index: int,
    context: str = "",
) -> list[ParticipantCandidate]:
    """Extract organizations explicitly listed in bidder/review/result tables.

    A supplier column by itself is insufficient: the header or nearby table
    heading must identify bidding, review, outcome, or award context.  Ranking
    fields are evidence of participation only and are never used to label a
    supplier as a winner or nonwinner.
    """
    result: list[ParticipantCandidate] = _direct_winner_key_values(
        rows, source_file=source_file, table_index=table_index, context=context,
    )
    columns: dict[str, int] = {}
    header_index: int | None = None
    for index, row in enumerate(rows[:10]):
        current = _participant_column_map(row)
        if "organization" in current:
            columns, header_index = current, index
            break
    if header_index is None:
        return result

    preamble = " ".join(" ".join(row) for row in rows[:header_index])
    header_text = " ".join(rows[header_index])
    evidence_context = _header_key(f"{context} {preamble} {header_text}")
    review_signal = any(_header_key(signal) in evidence_context for signal in PARTICIPANT_REVIEW_SIGNALS)
    context_signal = any(_header_key(signal) in evidence_context for signal in PARTICIPANT_CONTEXT_SIGNALS)
    explicit_result_columns = bool(
        {"outcome", "award_amount", "unsuccessful_reason"} & columns.keys()
    )
    if not (review_signal or context_signal or explicit_result_columns):
        return result

    package_hint = _explicit_package_hint([[context], *rows[:header_index]])
    org_column = columns["organization"]
    package_column = columns.get("package")
    current_package = package_hint
    header_width = len(rows[header_index])
    for row_index, raw_row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = [_clean(value) for value in raw_row]
        aligned_values = list(values)
        if (
            package_column is not None and package_column < len(values)
            and len(values) >= header_width and values[package_column]
        ):
            row_package = _participant_package_code(values[package_column])
            if row_package != "default":
                current_package = row_package
        status_columns = [
            columns[field] for field in ("unsuccessful_reason", "outcome", "qualification", "compliance")
            if field in columns
        ]
        # Some CCGP "未中标原因" tables put the package only in the first row;
        # later rows omit that cell and shift supplier/reason one column left.
        if (
            package_column == 0 and org_column == 1 and len(values) < header_width
            and len(values) >= 2 and any(column >= 2 for column in status_columns)
            and not _invalid_participant_name(values[0])
            and _invalid_participant_name(values[1])
        ):
            aligned_values = [""] * header_width
            aligned_values[org_column] = values[0]
            status_column = next(column for column in status_columns if column >= 2)
            aligned_values[status_column] = values[1]
            if current_package and package_column < header_width:
                aligned_values[package_column] = current_package
        if org_column >= len(aligned_values):
            continue
        organization = aligned_values[org_column].strip(" \t\r\n:：")
        normalized_org = _header_key(organization)
        if _invalid_participant_name(organization) or normalized_org in {
            _header_key(alias) for alias in PARTICIPANT_ORG_ALIASES
        }:
            continue

        def cell(field: str, row: list[str] = aligned_values) -> str:
            column = columns.get(field)
            return row[column] if column is not None and column < len(row) else ""

        raw_amount = cell("award_amount")
        amount_header = (
            rows[header_index][columns["award_amount"]]
            if "award_amount" in columns and columns["award_amount"] < len(rows[header_index])
            else ""
        )
        amount = _money(raw_amount, amount_header) if raw_amount else None
        outcome = _participant_outcome(aligned_values, columns, award_amount=amount)
        package = _participant_package_code(cell("package") or current_package or package_hint)
        evidence = " | ".join(value for value in values if value)
        if preamble:
            evidence = f"{_clean(preamble)} | {evidence}" if evidence else _clean(preamble)
        result.append(ParticipantCandidate(
            organization_name=organization,
            package_code=package,
            outcome=outcome,
            award_amount=amount if outcome == "winner" else None,
            source_file=source_file,
            source_location=f"table:{table_index}/row:{row_index}",
            source_evidence=evidence or organization,
            extraction_method="participant_table_header_mapping",
            confidence=0.8,
        ))
    return result


_NON_ITEM_ROW = re.compile(
    r"^(?:[一二三四五六七八九十百\d]+\s*[、)）]|"
    r"(?:联系人|联系方式|联系电话|联系地址|收费标准|采购代理机构)[：:]|"
    r"(?:中标情况|成交情况|评审专家名单|序号|包号|合计|总计|备注|说明)(?:$|[：:]))"
)
_REFERENCE_VALUE = re.compile(r"^(?:详?见|具体见|以).*(?:附件|文件|清单|响应|为准).*")


def _is_item_data_row(
    values: list[str], columns: dict[str, int], candidate: ItemCandidate
) -> bool:
    """Reject headings and narrative rows that happen to sit below an item header."""
    identity = candidate.product_name or candidate.category
    if not identity:
        return False
    if not _clean(identity) or _NON_ITEM_ROW.match(_clean(identity)):
        return False
    if _REFERENCE_VALUE.fullmatch(identity) or _header_key(identity) in {
        _header_key(alias) for field in ("product_name", "category")
        for alias in FIELD_ALIASES[field]
    }:
        return False

    # A real item row should populate at least two separately mapped cells (for
    # example name + quantity). This excludes merged section titles and contact
    # paragraphs, while allowing a sparse name/category row.
    populated_columns = {
        column for column in set(columns.values())
        if column < len(values) and values[column]
    }
    if len(populated_columns) < 2:
        return False

    # The same source cell must never be reused as both the item identity and its
    # supporting detail; this also guards against malformed merged-cell tables.
    identity_field = "product_name" if candidate.product_name else "category"
    identity_column = columns.get(identity_field)
    return identity_column is not None and bool(populated_columns - {identity_column})


def _html_tables(
    content: bytes, filename: str, warnings: list[str] | None = None,
) -> tuple[str, list[ItemCandidate], list[ParticipantCandidate]]:
    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    items: list[ItemCandidate] = []
    participants: list[ParticipantCandidate] = []
    reference_tables = []
    for table_index, table in enumerate(soup.find_all("table"), start=1):
        rows = [
            ["" if cell.find("table") else _clean(cell.get_text(" ", strip=True))
             for cell in row.find_all(["th", "td"], recursive=False)]
            for row in table.find_all("tr") if row.find_parent("table") is table
        ]
        headers = [(row, _column_map(row)) for row in rows if _item_header(_column_map(row))]
        if (headers and not table.find('table')
                and all(_requirement_header(row, columns) for row, columns in headers)):
            reference_tables.append(table)
            if warnings is not None:
                warnings.append(f'{filename}: table:{table_index} 为采购需求表，不抽取成交标的或送入模型')
            continue
        items.extend(parse_item_tables(rows, source_file=filename, table_index=table_index))
        context_parts = []
        caption = table.find("caption", recursive=False)
        if caption is not None:
            context_parts.append(caption.get_text(" ", strip=True))
        for attribute in ("summary", "aria-label", "title"):
            if table.get(attribute):
                context_parts.append(str(table.get(attribute)))
        # Only inspect immediately adjacent heading/paragraph siblings. A
        # document-wide previous heading can belong to an unrelated section.
        sibling = table.previous_sibling
        while sibling is not None:
            if isinstance(sibling, str) and not sibling.strip():
                sibling = sibling.previous_sibling
                continue
            name = getattr(sibling, "name", None)
            if name == "br":
                sibling = sibling.previous_sibling
                continue
            if isinstance(sibling, str):
                neighbor_text = _clean(sibling)
                if neighbor_text and len(neighbor_text) <= 120:
                    context_parts.append(neighbor_text)
                break
            if name in {"p", "h1", "h2", "h3", "h4", "strong", "b"}:
                neighbor_text = _clean(sibling.get_text(" ", strip=True))
                if neighbor_text and len(neighbor_text) <= 120:
                    context_parts.append(neighbor_text)
                sibling = sibling.previous_sibling
                break
            break
        # Result pages sometimes put each package number in the first cell of
        # an outer "包号 / 供货明细" table and nest the bidder table beside it.
        for containing_row in table.find_parents("tr"):
            containing_table = containing_row.find_parent("table")
            if containing_table is None:
                continue
            header_row = next(
                (row for row in containing_table.find_all("tr")
                 if row.find_parent("table") is containing_table),
                None,
            )
            if header_row is None:
                continue
            header_cells = header_row.find_all(["th", "td"], recursive=False)
            package_header = any(
                _header_key(cell.get_text(" ", strip=True)) in {
                    _header_key(alias) for alias in PARTICIPANT_PACKAGE_ALIASES
                }
                for cell in header_cells
            )
            if not package_header:
                continue
            row_cells = containing_row.find_all(["th", "td"], recursive=False)
            child_index = next(
                (index for index, cell in enumerate(row_cells)
                 if cell is table.parent or cell.find("table") is table),
                None,
            )
            if child_index:
                package_value = row_cells[child_index - 1].get_text(" ", strip=True)
                if package_value:
                    context_parts.append(f"包号 {package_value}")
            break
        participants.extend(parse_participant_tables(
            rows, source_file=filename, table_index=table_index,
            context=" ".join(context_parts),
        ))
    for table in reference_tables:
        table.replace_with('[采购需求表已排除，原文保留在来源文件]')
    # Keep block and cell boundaries for metadata; a space-flattened document
    # cannot distinguish the buyer from the next row's administrative region.
    for tag in soup.find_all(["p", "div", "tr", "table", "section", "br",
                              "h1", "h2", "h3", "h4", "li", "title"]):
        tag.insert_before("\n")
        tag.insert_after("\n")
    for cell in soup.find_all(["td", "th"]):
        # Cells frequently contain <p>/<div> wrappers; these are not new rows.
        if not cell.find("table"):
            value = _clean(cell.get_text(" ", strip=True))
            cell.clear()
            cell.append(value)
        cell.insert_after("\t")
    text = "\n".join(
        re.sub(r"[^\S\t]+", " ", line).strip(" ")
        for line in soup.get_text("").splitlines() if line.strip()
    )
    return text, items, participants


def _docx_tables(
    content: bytes, filename: str,
) -> tuple[str, list[ItemCandidate], list[ParticipantCandidate]]:
    document = Document(io.BytesIO(content))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    items: list[ItemCandidate] = []
    participants: list[ParticipantCandidate] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        chunks.extend(" ".join(_clean(cell) for cell in row) for row in rows)
        items.extend(parse_item_tables(rows, source_file=filename, table_index=table_index))
        participants.extend(parse_participant_tables(
            rows, source_file=filename, table_index=table_index,
        ))
    return _clean(" ".join(chunks)), items, participants


def _xlsx_tables(
    content: bytes, filename: str,
) -> tuple[str, list[ItemCandidate], list[ParticipantCandidate]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    chunks: list[str] = []
    items: list[ItemCandidate] = []
    participants: list[ParticipantCandidate] = []
    table_index = 0
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows = [[_clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(row)]
        chunks.extend(" ".join(row) for row in rows)
        parsed = parse_item_tables(rows, source_file=filename, table_index=sheet_index)
        if parsed:
            table_index += 1
            items.extend(parsed)
        participants.extend(parse_participant_tables(
            rows, source_file=filename, table_index=sheet_index, context=sheet.title,
        ))
    workbook.close()
    return _clean(" ".join(chunks)), items, participants


def _xls_tables(
    content: bytes, filename: str,
) -> tuple[str, list[ItemCandidate], list[ParticipantCandidate]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    chunks: list[str] = []
    items: list[ItemCandidate] = []
    participants: list[ParticipantCandidate] = []
    try:
        for index, sheet in enumerate(workbook.sheets(), start=1):
            # Keep numeric zero (e.g. a free item).
            rows = [[str(cell.value) if cell.ctype == xlrd.XL_CELL_NUMBER
                     else _clean(cell.value) for cell in row] for row in sheet.get_rows()]
            chunks.extend('\t'.join(row) for row in rows if any(row))
            parsed = parse_item_tables(rows, source_file=filename, table_index=index)
            for item in parsed:
                item.extraction_method = 'xls_table_header_mapping'
            items.extend(parsed)
            participants.extend(parse_participant_tables(
                rows, source_file=filename, table_index=index, context=sheet.name,
            ))
        return '\n'.join(chunks), items, participants
    finally:
        workbook.release_resources()


def parser_capabilities() -> dict[str, bool]:
    return {
        "pdf_text": True,
        "pdf_tables": importlib.util.find_spec("pdfplumber") is not None,
        "pdf_render": importlib.util.find_spec("pypdfium2") is not None,
        "image_ocr": (importlib.util.find_spec('rapidocr') is not None
                      and importlib.util.find_spec('onnxruntime') is not None
                      if settings.ocr_engine == 'rapidocr' else shutil.which('tesseract') is not None),
        "archive_7z": importlib.util.find_spec('py7zr') is not None,
        "legacy_doc": libreoffice_executable() is not None,
        "legacy_xls": importlib.util.find_spec("xlrd") is not None,
    }


def _ocr_image(
    content: bytes, filename: str, *, language: str, timeout: int,
    engine: Callable[[bytes, str], str] | None = None,
) -> tuple[str, list[str]]:
    if engine is not None:
        return engine(content, filename), []
    if settings.ocr_engine == 'rapidocr':
        try:
            from app.local_ocr import recognize
            text, warnings = recognize(content)
            return text, [f'{filename}: {message}' for message in warnings]
        except ImportError:
            return '', [f'{filename}: RapidOCR 未安装，请安装后端 ocr 依赖']
        except Exception as exc:  # noqa: BLE001
            return '', [f'{filename}: 本地 OCR 失败（{type(exc).__name__}）']
    executable = shutil.which("tesseract")
    if executable is None:
        return "", [f"{filename}: OCR 已启用但未找到 Tesseract，可安装并配置 chi_sim 语言包"]
    if not re.fullmatch(r"[A-Za-z0-9_+\-]+", language):
        return "", [f"{filename}: OCR 语言参数无效"]
    with tempfile.TemporaryDirectory(prefix="bidintel-ocr-") as directory:
        path = Path(directory) / ("page" + (Path(filename).suffix or ".png"))
        path.write_bytes(content)
        try:
            result = subprocess.run(
                [executable, str(path), "stdout", "-l", language],
                capture_output=True, timeout=max(1, min(timeout, 120)), check=False,
            )
        except subprocess.TimeoutExpired:
            return "", [f"{filename}: OCR 超时，已跳过该图像"]
        if result.returncode:
            return "", [f"{filename}: Tesseract OCR 失败，请检查图像与语言包"]
        return result.stdout.decode("utf-8", errors="replace"), []


def _pdf_text(
    content: bytes, filename: str, *, ocr_enabled: bool = False,
    ocr_language: str = "chi_sim+eng", ocr_timeout_seconds: int = 30,
    ocr_engine: Callable[[bytes, str], str] | None = None,
) -> tuple[str, list[ItemCandidate], list[ParticipantCandidate], list[str]]:
    reader = PdfReader(io.BytesIO(content))
    page_text: list[str] = []
    items: list[ItemCandidate] = []
    participants: list[ParticipantCandidate] = []
    warnings: list[str] = []
    scan_pages = []
    page_limit = max(1, settings.pdf_max_pages)
    ocr_limit = max(1, settings.pdf_max_ocr_pages)
    for page_number, page in enumerate(reader.pages[:page_limit], start=1):
        text = page.extract_text() or ""
        if text.strip():
            page_text.append(f"[page:{page_number}] {text}")
        else:
            scan_pages.append(page_number)
    if len(reader.pages) > page_limit:
        warnings.append(f'{filename}: PDF 超过 {page_limit} 页，只解析前 {page_limit} 页')
    if importlib.util.find_spec("pdfplumber"):
        try:
            with importlib.import_module("pdfplumber").open(io.BytesIO(content)) as pdf:
                for page_number, page in enumerate(pdf.pages[:page_limit], start=1):
                    if page_number in scan_pages:
                        page.close()
                        continue
                    for table_index, table in enumerate(page.extract_tables() or [], start=1):
                        parsed = parse_item_tables(table, source_file=filename, table_index=table_index)
                        for item in parsed:
                            item.source_location = f"page:{page_number}/{item.source_location}"
                            item.extraction_method = "pdfplumber_table_header_mapping"
                        items.extend(parsed)
                        participants.extend(parse_participant_tables(
                            table, source_file=filename, table_index=table_index,
                        ))
                        for participant in participants:
                            if participant.source_file == filename and participant.source_location.startswith(
                                f"table:{table_index}/"
                            ):
                                participant.source_location = (
                                    f"page:{page_number}/{participant.source_location}"
                                )
                                participant.extraction_method = "pdfplumber_participant_table"
                    page.close()
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"{filename}: PDF 表格解析失败（{type(exc).__name__}），保留文本结果")
    else:
        warnings.append(f"{filename}: 缺少 pdfplumber，请重装后端依赖；本次 PDF 仅提取文本")
    if scan_pages:
        if not ocr_enabled:
            warnings.append(f"{filename}: 第 {','.join(map(str, scan_pages))} 页无可提取文本，OCR 未启用")
        elif not importlib.util.find_spec("pypdfium2"):
            warnings.append(f"{filename}: 缺少 pypdfium2 渲染组件，请重装后端依赖")
        else:
            renderer = importlib.import_module("pypdfium2").PdfDocument(content)
            try:
                for page_number in scan_pages[:ocr_limit]:
                    page = renderer[page_number - 1]
                    bitmap = page.render(scale=2)
                    buffer = io.BytesIO()
                    try:
                        bitmap.to_pil().save(buffer, format="PNG")
                        text, ocr_warnings = _ocr_image(
                            buffer.getvalue(), f"{filename}.page{page_number}.png",
                            language=ocr_language, timeout=ocr_timeout_seconds, engine=ocr_engine,
                        )
                        if text.strip():
                            page_text.append(f"[page:{page_number}/ocr] {text}")
                            parsed = parse_item_tables([line.split('\t') for line in text.splitlines()],
                                                       source_file=filename, table_index=1)
                            for item in parsed:
                                item.source_location = f'page:{page_number}/ocr/{item.source_location}'
                                item.extraction_method = 'local_ocr_table'
                            items.extend(parsed)
                        warnings.extend(ocr_warnings)
                    finally:
                        bitmap.close()
                        page.close()
                if len(scan_pages) > ocr_limit:
                    warnings.append(f'{filename}: 每份 PDF 最多 OCR {ocr_limit} 页，其余扫描页未处理')
            finally:
                renderer.close()
    return _clean(" ".join(page_text)), items, participants, warnings


def parse_document_with_participants(
    document: SourceDocument, *, ocr_enabled: bool = False,
    ocr_language: str = "chi_sim+eng", ocr_timeout_seconds: int = 30,
    ocr_engine: Callable[[bytes, str], str] | None = None,
) -> tuple[str, list[ItemCandidate], list[ParticipantCandidate], list[str]]:
    suffix = PurePosixPath(document.filename.replace("\\", "/")).suffix.lower()
    warnings: list[str] = []
    try:
        detected = inspect_content(document.content)
        warnings.extend(f"{document.filename}: {message}" for message in detected.warnings)
        if detected.error:
            return "", [], [], warnings + [f"{document.filename}: {detected.error}"]
        original_suffix = suffix
        if detected.format:
            suffix = "." + detected.format
            aliases = {".htm": ".html", ".jpeg": ".jpg", ".tif": ".tiff"}
            if aliases.get(original_suffix, original_suffix) != suffix:
                label = {"html": "HTML", "docx": "OOXML DOCX", "xlsx": "OOXML XLSX"}.get(
                    detected.format, detected.format.upper())
                warnings.append(f"{document.filename}: 实际为 {label}，已按内容解析（原扩展名 {original_suffix or '无'}）")
        content = detected.content
        if suffix in {".html", ".htm"}:
            text, items, participants = _html_tables(content, document.filename, warnings)
        elif suffix == ".xls":
            text, items, participants = _xls_tables(content, document.filename)
        elif suffix in {".doc", ".rtf"}:
            text, items, participants = _docx_tables(convert_doc(content), document.filename)
            for item in items:
                item.extraction_method = 'doc_converted_table_header_mapping'
            if text and not items:
                warnings.append(f"{document.filename}: DOC 文本已读取，未映射出标的表格，请核验布局或使用模型")
        elif suffix == ".docx":
            text, items, participants = _docx_tables(content, document.filename)
        elif suffix == ".xlsx":
            text, items, participants = _xlsx_tables(content, document.filename)
        elif suffix == ".pdf":
            text, items, participants, pdf_warnings = _pdf_text(
                content, document.filename, ocr_enabled=ocr_enabled,
                ocr_language=ocr_language, ocr_timeout_seconds=ocr_timeout_seconds,
                ocr_engine=ocr_engine,
            )
            warnings.extend(pdf_warnings)
            if text and not items:
                warnings.append(
                    f"{document.filename}: PDF 文本已读取；复杂表格/扫描页需 OCR 或模型解析后核验"
                )
        elif suffix == ".txt":
            text, items = content.decode("utf-8", errors="replace"), []
            participants = parse_participant_tables(
                [line.split("\t") for line in text.splitlines()],
                source_file=document.filename, table_index=1,
            )
        elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
            if not ocr_enabled:
                return "", [], [], warnings + [f"{document.filename}: 图像附件需要 OCR，当前未启用"]
            text, ocr_warnings = _ocr_image(
                content, document.filename, language=ocr_language,
                timeout=ocr_timeout_seconds, engine=ocr_engine,
            )
            warnings.extend(ocr_warnings)
            items = parse_item_tables([line.split('\t') for line in text.splitlines()],
                                      source_file=document.filename, table_index=1)
            participants = parse_participant_tables(
                [line.split('\t') for line in text.splitlines()],
                source_file=document.filename, table_index=1,
            )
            for item in items:
                item.extraction_method = 'local_ocr_table'
        else:
            return "", [], [], warnings + [f"暂不支持附件格式：{document.filename}"]
    except DocumentConversionError as exc:
        return "", [], [], warnings + [f"解析失败：{document.filename}（{exc}）"]
    except Exception as exc:  # noqa: BLE001 - isolate individual corrupt attachments
        return "", [], [], warnings + [f"解析失败：{document.filename}（{type(exc).__name__}）"]
    if not text.strip() and not items and not participants and not warnings:
        warnings.append(f"{document.filename}: 未解析出文本或表格，请检查空文件、加密或扫描内容")
    return text.strip(), items, participants, warnings


def parse_document(
    document: SourceDocument, *, ocr_enabled: bool = False,
    ocr_language: str = "chi_sim+eng", ocr_timeout_seconds: int = 30,
    ocr_engine: Callable[[bytes, str], str] | None = None,
) -> tuple[str, list[ItemCandidate], list[str]]:
    """Backward-compatible parser API; participant-aware callers use details API."""
    text, items, _, warnings = parse_document_with_participants(
        document, ocr_enabled=ocr_enabled, ocr_language=ocr_language,
        ocr_timeout_seconds=ocr_timeout_seconds, ocr_engine=ocr_engine,
    )
    return text, items, warnings


def _zip_member_name(info: zipfile.ZipInfo, warnings: list[str]) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    raw = info.orig_filename.encode('cp437')
    # Info-ZIP Unicode Path overrides a legacy name only when its CRC matches.
    extra = info.extra
    while len(extra) >= 4:
        kind, size = struct.unpack_from('<HH', extra)
        payload, extra = extra[4:4 + size], extra[4 + size:]
        if kind == 0x7075 and len(payload) >= 5 and payload[0] == 1:
            if struct.unpack_from('<I', payload, 1)[0] == zlib.crc32(raw):
                try:
                    return payload[5:].decode('utf-8')
                except UnicodeDecodeError:
                    pass
            warnings.append(f"ZIP Unicode 文件名校验失败，尝试原始编码：{info.filename}")
    for encoding in ('utf-8', 'gb18030'):
        try:
            name = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if name != info.filename:
            warnings.append(f"ZIP 文件名按 {encoding} 解码：{name}")
        return name
    warnings.append(f"ZIP 文件名编码无法确定，保留原名，请核验附件归属：{info.filename}")
    return info.filename


def expand_uploads(files: list[SourceDocument]) -> tuple[list[SourceDocument], list[str]]:
    """Expand nested ZIPs in memory with count/size/depth guards."""
    expanded: list[SourceDocument] = []
    warnings: list[str] = []
    total_bytes = 0

    def visit(document: SourceDocument, depth: int) -> None:
        nonlocal total_bytes
        suffix = PurePosixPath(document.filename.replace("\\", "/")).suffix.lower()
        detected = inspect_content(document.content)
        warnings.extend(f"{document.filename}: {message}" for message in detected.warnings)
        if detected.error:
            warnings.append(f"{document.filename}: {detected.error}")
            return
        if detected.content is not document.content:
            document = SourceDocument(document.filename, detected.content)
        is_zip = (detected.format == 'zip' and suffix not in {'.gbq7'}
                  or suffix == '.zip' and detected.format is None)
        if is_zip and suffix != '.zip':
            warnings.append(f"{document.filename}: 按实际 ZIP 内容展开，保留原名")
        if not is_zip and suffix == '.zip' and detected.format:
            warnings.append(f"{document.filename}: 扩展名 ZIP 与实际 {detected.format} 不符，按内容解析")
        if not is_zip:
            total_bytes += len(document.content)
            if total_bytes > MAX_EXPANDED_BYTES:
                raise ValueError("解压后数据超过 200 MB 限制")
            expanded.append(document)
            return
        if depth >= MAX_ARCHIVE_DEPTH:
            warnings.append(f"压缩包嵌套超过 {MAX_ARCHIVE_DEPTH} 层：{document.filename}")
            return
        try:
            with zipfile.ZipFile(io.BytesIO(document.content)) as archive:
                members = [info for info in archive.infolist() if not info.is_dir()]
                if len(expanded) + len(members) > MAX_ARCHIVE_FILES:
                    raise ValueError("压缩包文件数量超过限制")
                seen_names: set[str] = set()
                for info in members:
                    name = _zip_member_name(info, warnings).replace('\\', '/')
                    if name in seen_names:
                        warnings.append(f"ZIP 解码后文件名重复，跳过后续同名文件：{document.filename}!/{name}")
                        continue
                    seen_names.add(name)
                    # Do not write archive paths to disk; names are kept only as provenance.
                    if info.file_size > MAX_EXPANDED_BYTES:
                        warnings.append(f"跳过超大压缩文件：{name}")
                        continue
                    try:
                        content = archive.read(info)
                    except (zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError):
                        warnings.append(f"ZIP 成员读取失败（损坏、加密或压缩算法不支持）：{name}")
                        continue
                    nested = SourceDocument(
                        filename=f"{document.filename}!/{name}",
                        content=content,
                    )
                    visit(nested, depth + 1)
        except (zipfile.BadZipFile, UnicodeDecodeError):
            warnings.append(f"压缩包损坏或格式无效：{document.filename}")

    for uploaded in files:
        visit(uploaded, 0)
    return expanded, warnings


METADATA_LABELS = {
    "project_name": ("采购项目名称", "项目名称"),
    "project_number": ("采购项目编号", "项目编号", "采购计划编号"),
    "procurement_unit": ("采购单位名称", "采购单位", "采购人名称", "采购人", "采购机构"),
    "project_budget": ("采购预算金额", "采购预算", "预算金额", "项目预算金额", "项目预算"),
    "announced_total_award": (
        "总中标金额", "总成交金额", "总中标/成交金额", "中标（成交）金额",
        "中标/成交金额", "中标金额", "成交金额", "项目成交金额",
    ),
}


def extract_metadata(text: str) -> NoticeMetadata:
    """Prefer exact table label/value cells, then bounded colon-labelled lines."""
    label_fields = {label: field for field, labels in METADATA_LABELS.items() for label in labels}
    cell_fields = {_header_key(label): field for label, field in label_fields.items()}
    values: dict[str, object] = {}

    def accept(field: str, raw: str, label: str) -> None:
        value = _clean(raw).strip(" ：:，,。")
        if not value or field in values or _REFERENCE_VALUE.fullmatch(value):
            return
        if field in {"project_budget", "announced_total_award"}:
            amount = _money(value, label)
            if amount is not None:
                values[field] = amount
        elif field == "project_number":
            # Codes may contain parentheses but cannot include a title suffix.
            code = re.split(r"[）)](?=.*(?:公告|采购项目))", value, maxsplit=1)[0]
            if len(code) <= 120:
                values[field] = code
        elif len(value) <= 160:
            values[field] = value

    lines = text.splitlines()
    for line in lines:
        # HTML extraction preserves empty cells and tabs, including 4-column
        # rows containing two key/value pairs. Never join adjacent row values.
        cells = line.rstrip("\t ").split("\t")
        for index in range(0, len(cells) - 1, 2):
            label, value = cells[index:index + 2]
            field = cell_fields.get(_header_key(label))
            if field and _header_key(value) not in cell_fields:
                accept(field, value, label)

    pattern = "|".join(re.escape(label) for label in sorted(label_fields, key=len, reverse=True))
    labelled = re.compile(rf"(?P<label>{pattern})\s*[：:]\s*")
    for line in lines:
        # A label on its own (e.g. 项目名称) is not a value. No whitespace-only
        # fallback: structural table rows were handled above.
        matches = list(labelled.finditer(line))
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(line)
            raw = line[match.end():end]
            raw = re.split(r"[\t；;。]|\s+[\u4e00-\u9fff]{2,16}\s*[：:]|"
                           r"\s*[一二三四五六七八九十]+\s*[、）)]|"
                           r"\s+\d+\s*[、）)]", raw, maxsplit=1)[0]
            accept(label_fields[match["label"]], raw, match["label"])
    return NoticeMetadata(**values)
