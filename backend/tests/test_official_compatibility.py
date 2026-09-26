"""Regressions for formats, roles and cross-file ambiguity in the official intake."""
import gzip
import io
import os
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from docx import Document
from filelock import Timeout
from reportlab.pdfgen.canvas import Canvas

from app import ingestion, legacy_documents
from app.config import Settings
from app.file_formats import inspect_content
from app.ingestion import _deduplicate, extract_notice
from app.parsers import SourceDocument, expand_uploads, parse_document, parse_item_tables
from app.schemas import ItemCandidate, NoticeMetadata


def docx_bytes():
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph('quotation')
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize('filename', ['opaque', 'wrong.docx'])
def test_pdf_content_overrides_extension_and_keeps_source(filename):
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(10, 10, 'quotation')
    canvas.save()
    text, _, warnings = parse_document(SourceDocument(filename, buffer.getvalue()))
    assert 'quotation' in text
    assert any('实际为 PDF' in row and filename in row for row in warnings)


def test_extensionless_image_reaches_ocr():
    calls = []

    def ocr(content, filename):
        calls.append(filename)
        return '已读取图片'

    text, _, warnings = parse_document(SourceDocument('opaque', b'\x89PNG\r\n\x1a\n'),
                                      ocr_enabled=True, ocr_engine=ocr)
    assert text == '已读取图片' and calls == ['opaque']
    assert '实际为 PNG' in warnings[0]


def test_docx_named_zip_stays_one_document_and_legacy_doc_is_detected():
    rows, _ = expand_uploads([SourceDocument('wrong.zip', docx_bytes())])
    assert len(rows) == 1 and rows[0].filename == 'wrong.zip'
    assert parse_document(rows[0])[0] == 'quotation'
    payload = (Path(__file__).parent / 'fixtures/legacy-quotation.doc').read_bytes()
    assert inspect_content(payload).format == 'doc'


def test_zip_with_pdf_member_signature_is_not_misdetected_as_pdf(tmp_path):
    pdf = io.BytesIO()
    canvas = Canvas(pdf)
    canvas.drawString(10, 10, 'quotation')
    canvas.save()
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, 'w') as output:
        output.writestr('报价.pdf', pdf.getvalue())
    assert inspect_content(archive.getvalue()).format == 'zip'
    outer = tmp_path / 'outer.zip'
    outer.write_bytes(archive.getvalue())
    documents, _ = expand_uploads([SourceDocument('outer.zip', outer.read_bytes())])
    assert documents[0].filename.endswith('outer.zip!/报价.pdf')
    assert 'quotation' in parse_document(documents[0])[0]


@pytest.mark.parametrize('payload', [
    '<html><title>下载</title><body>请输入验证码 再次下载请刷新页面</body></html>',
    '<html><body>系统正在维护中</body></html>',
    '<html><title>系统限制</title><body>访问受限</body></html>',
    '{"successFul": false, "message": "download failed"}',
])
def test_download_error_isolated_before_model(payload, monkeypatch):
    calls = []

    def model(**kwargs):
        calls.append(kwargs['filename'])
        return NoticeMetadata(), [], [], []

    monkeypatch.setattr(ingestion, 'extract_unstructured_items', model)
    result = extract_notice([
        SourceDocument('a.html', b'<p>Notice</p>'),
        SourceDocument('a.zip!/bad.pdf', gzip.compress(payload.encode())),
    ], extraction_mode='hybrid', model_settings=Settings(
        _env_file=None, model_base_url='https://example.invalid', model_api_key='stub', model_name='stub',
    ))
    assert calls == ['a.html']
    assert any('源附件不可用' in row and 'bad.pdf' in row for row in result.warnings)
    assert not result.items


def test_legitimate_captcha_requirement_is_not_download_error():
    payload = '<html><title>采购公告</title><p>技术要求：用户登录请输入验证码。</p></html>'.encode()
    assert inspect_content(payload).error is None


def test_error_only_upload_explains_unavailable_source():
    with pytest.raises(ValueError, match='源附件不可用'):
        extract_notice([SourceDocument('bad.pdf', b'{"success": false}')], extraction_mode='rules')


def test_conversion_cache_reuses_across_threads_and_rebuilds_bad_cache(tmp_path, monkeypatch):
    executable = tmp_path / 'soffice.exe'
    executable.write_bytes(b'fake')
    monkeypatch.setattr(legacy_documents, 'libreoffice_executable', lambda: str(executable))
    calls = []
    converted = docx_bytes()

    def convert(content, *args):
        calls.append(content)
        return converted

    monkeypatch.setattr(legacy_documents, '_convert_uncached', convert)
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(legacy_documents.convert_doc, [b'one', b'one'])) == [converted, converted]
    assert calls == [b'one']
    root = legacy_documents.settings.resolved_database_path.parent / 'document-cache'
    next(root.glob('*.docx')).write_bytes(b'corrupt')
    assert legacy_documents.convert_doc(b'one') == converted
    assert calls == [b'one', b'one']
    legacy_documents.convert_doc(b'two')
    stamp = executable.stat()
    os.utime(executable, ns=(stamp.st_atime_ns, stamp.st_mtime_ns + 1_000_000_000))
    legacy_documents.convert_doc(b'one')
    assert calls == [b'one', b'one', b'two', b'one']


def test_failed_conversion_is_not_cached_and_queue_timeout_is_clear(monkeypatch):
    monkeypatch.setattr(legacy_documents, 'libreoffice_executable', lambda: 'soffice')

    def fail(*args):
        raise legacy_documents.DocumentConversionError('failed')

    monkeypatch.setattr(legacy_documents, '_convert_uncached', fail)
    with pytest.raises(legacy_documents.DocumentConversionError):
        legacy_documents.convert_doc(b'one')
    root = legacy_documents.settings.resolved_database_path.parent / 'document-cache'
    assert not list(root.glob('*.docx'))

    def busy(*args, **kwargs):
        raise Timeout('converter.lock')

    monkeypatch.setattr(legacy_documents, 'FileLock', busy)
    with pytest.raises(legacy_documents.DocumentConversionError, match='队列等待'):
        legacy_documents.convert_doc(b'one')


def test_requirement_table_removed_from_rules_and_model_text():
    html = '''<table><tr><td>包号</td><td>标的名称</td><td>简要技术要求</td><td>数量</td></tr>
    <tr><td>1</td><td>需求设备</td><td>详见招标文件</td><td>1</td></tr></table>
    <table><tr><td>名称</td><td>数量</td><td>单价</td></tr>
    <tr><td>成交设备</td><td>1</td><td>100</td></tr></table>'''
    text, items, warnings = parse_document(SourceDocument('a.html', html.encode()))
    assert [item.product_name for item in items] == ['成交设备']
    assert '需求设备' not in text and '成交设备' in text
    assert any('采购需求表' in row for row in warnings)
    # A requirement-shaped outer table must not remove a nested actual award.
    nested = html.replace('</table>\n    <table>', '<tr><td><table>', 1)
    nested += '</td></tr></table>'
    text, items, _ = parse_document(SourceDocument('nested.html', nested.encode()))
    assert '成交设备' in text
    assert [item.product_name for item in items] == ['成交设备']


def test_reference_role_uses_leaf_name_and_keeps_real_quotes(monkeypatch):
    files = [SourceDocument(name, b'<p>Notice</p>') for name in [
        'a.html', '采购文件集.zip!/采购文件.docx', '采购文件集.zip!/成交报价明细.docx',
    ]]
    calls = []

    def parse(document, **kwargs):
        calls.append(document.filename)
        return '项目名称：正确项目', [], []

    monkeypatch.setattr(ingestion, 'parse_document', parse)
    result = extract_notice(files, extraction_mode='rules', model_settings=Settings(_env_file=None))
    assert calls == [files[0].filename, files[2].filename]
    assert result.source_files == [row.filename for row in files]
    assert any('参考采购材料' in row for row in result.warnings)


def test_quotation_aliases_and_explicit_package_not_row_number():
    rows = [['货物类供应商报价成交明细'],
            ['序号', '品目分类', '货物名称', '品牌', '规格说明', '数量', '单价', '报价'],
            ['包1', '设备', '打印机', 'A', 'X1', '1.0', '100.0', '100.0'],
            ['2', '设备', '扫描仪', 'B', 'X2', '1', '200', '200']]
    items = parse_item_tables(rows, source_file='quote.xls', table_index=1)
    assert [(row.package_code, row.model, row.total_price) for row in items] == [
        ('包1', 'X1', 100), ('default', 'X2', 200),
    ]
    assert parse_item_tables(rows[1:], source_file='unknown.xls', table_index=1)[0].total_price is None


def item(source, **changes):
    values = {'product_name': '打印机', 'model': 'X1', 'quantity': '1', 'unit_price': '100',
                  'source_file': source, 'source_location': 'row:1', 'source_evidence': '原文'}
    values.update(changes)
    return ItemCandidate(**values)


def test_cross_file_decimal_matching_preserves_sources_and_completes_fields():
    warnings = []
    rows = _deduplicate([item('a.html'), item('quote.xls', quantity='1.0', unit_price='100.00',
                                             total_price='100', package_code='包1')], warnings)
    assert len(rows) == 1 and rows[0].total_price == 100 and rows[0].package_code == '包1'
    assert all(source in rows[0].source_evidence for source in ['a.html', 'quote.xls'])
    assert any('重复候选合并' in row for row in warnings)


@pytest.mark.parametrize('change', [{'package_code': '包2'}, {'brand': '另一品牌'},
                                     {'model': 'X2'}, {'quantity_unit': '套'}])
def test_cross_file_conflicts_preserved(change):
    left = item('a.html', package_code='包1', brand='A', quantity_unit='台')
    right = item('quote.xls', package_code='包1', brand='A', quantity_unit='台').model_copy(update=change)
    assert len(_deduplicate([left, right])) == 2


@pytest.mark.parametrize('reverse', [False, True])
def test_missing_package_cannot_bridge_two_packages(reverse):
    rows = [item('a.html', package_code='包1'), item('b.xls'), item('c.pdf', package_code='包2')]
    warnings = []
    output = _deduplicate(list(reversed(rows)) if reverse else rows, warnings)
    assert len(output) == 3 and any('无法唯一对齐' in row for row in warnings)


def test_same_file_repeats_and_sparse_rows_are_not_silently_deleted():
    assert len(_deduplicate([item('a.html'), item('a.html', source_location='row:2'), item('b.xls')])) == 3
    assert len(_deduplicate([item('a.html', model=None, unit_price=None),
                             item('b.xls', model=None, unit_price=None)])) == 2


def test_extensionless_zip_expands_with_original_path():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        archive.writestr('报价.txt', '真实报价')
    rows, warnings = expand_uploads([SourceDocument('opaque', buffer.getvalue())])
    assert [row.filename for row in rows] == ['opaque!/报价.txt']
    assert any('实际 ZIP' in row for row in warnings)


def test_unfilled_quote_template_is_not_sent_to_model(monkeypatch):
    blank = '投标人名称： 合计 备注：无 { 供应商响应 } {供应商响应}'
    filled = blank.replace('名称： 合计', '名称：某公司 合计')
    assert ingestion._unfilled_template(blank, [item('template.pdf')])
    assert not ingestion._unfilled_template(filled, [item('filled.pdf')])
    calls = []

    def model(**kwargs):
        calls.append(kwargs['filename'])
        return NoticeMetadata(), [], [], []

    monkeypatch.setattr(ingestion, 'extract_unstructured_items', model)
    result = extract_notice([
        SourceDocument('a.html', b'<p>notice</p>'),
        SourceDocument('quotation.txt', blank.encode()),
    ], extraction_mode='hybrid', model_settings=Settings(
        _env_file=None, model_base_url='https://example.invalid', model_api_key='stub', model_name='stub',
    ))
    assert calls == ['a.html']
    assert any('未填写模板' in row for row in result.warnings)
