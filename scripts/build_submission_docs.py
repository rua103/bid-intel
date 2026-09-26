from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "drafts"
GREEN = "176B65"
INK = "203235"
MUTED = "607174"
PALE = "EAF3F1"
GRID = "D5DEDD"
FONT = "Microsoft YaHei"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), color)


def set_run_font(run, size: float, *, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)


def style_document(doc: Document, title: str):
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor.from_string(INK if name == "Title" else GREEN)
        style.paragraph_format.space_before = Pt(9 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("招采数据分析引擎  ·  提交材料草稿")
    set_run_font(run, 8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("赛题五  |  工作草稿  |  "), 8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.subject = "赛题五项目提交材料草稿"


def add_cover(doc: Document, title: str, subtitle: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("赛题五  /  招采数据智能分析引擎"), 11, bold=True, color=GREEN)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(title), 25, bold=True, color="000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run(subtitle), 12, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("版本 0.1  ·  2026-09-26  ·  提交前草稿"), 9, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("评审口径："), 10, bold=True, color=GREEN)
    set_run_font(p.add_run("当前没有人工 Gold，也没有官方查询基准。官方数据上的候选、OCR 提示和吞吐记录只说明接入管线运行过，不能解释为准确率或评分结果。"), 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_para(doc: Document, text: str, *, lead: str | None = None):
    p = doc.add_paragraph()
    if lead and text.startswith(lead):
        set_run_font(p.add_run(lead), 10.3, bold=True, color=INK)
        set_run_font(p.add_run(text[len(lead):]), 10.3)
    else:
        set_run_font(p.add_run(text), 10.3)
    return p


def add_bullets(doc: Document, rows: list[str]):
    for text in rows:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(text), 10)


def add_numbered(doc: Document, rows: list[str]):
    for text in rows:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(text), 10)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, GREEN)
        set_cell_border(cell, GRID)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_run_font(cell.paragraphs[0].add_run(header), 9, bold=True, color="FFFFFF")
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if index % 2 == 1:
                set_cell_shading(cells[i], "F5F8F7")
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(3)
            set_run_font(cells[i].paragraphs[0].add_run(value), 8.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, 17 if level == 1 else 12.5, bold=True, color=GREEN)
    return p


def build_system_design() -> Path:
    doc = Document()
    style_document(doc, "招采数据智能分析引擎系统设计说明")
    add_cover(doc, "系统设计说明", "采购公告解析、关系建模与可核验查询")

    add_heading(doc, "一 项目目标与范围")
    add_para(doc, "系统接收采购公告正文和附件，提取七项标的字段及采购单位、投标主体与中标关系，将原文证据随结果入库，再提供采购单位供应商、高频投标主体、共同投标和共同合作单位等查询。系统把自动输出视为候选，留有来源文件、位置与原文证据供人工核验。")
    add_table(doc, ["输入", "处理", "输出"], [
        ["HTML、DOC/DOCX、XLS/XLSX、可读 PDF、ZIP/RAR/7z 及图片", "内容识别、归档展开、文本/表格解析、可选 OCR 与模型抽取", "采购单位、项目、采购包、投标主体、中标记录、标的候选及告警"],
    ])

    add_heading(doc, "二 业务与验收要求")
    add_bullets(doc, [
        "每个字段保留来源文件及位置；未披露值保持空，不用常识补全。",
        "区分投标参与、明确未中标、明确中标和结果未知；排名分数不能单独推断中标状态。",
        "附件解析问题要可见：未匹配、内容错误页、不支持格式、OCR 无可信文本及解析异常都应提示。",
        "查询按采购包建模；频次与金额定义以赛题正式字段说明和基准样例为最终口径。",
        "原始官方材料不提交到代码仓库；运行数据库、模型凭据、演示账号留在本地忽略目录。",
    ])

    add_heading(doc, "三 总体架构")
    add_para(doc, "前端使用 Vue 3/Vite；后端使用 FastAPI，按解析、抽取、持久化、关系查询及任务 API 分层；关系数据默认写入 SQLite。Neo4j 导出和查询代码独立于当前 API 主链路。")
    add_table(doc, ["层", "实现", "边界"], [
        ["交互层", "分析台、批处理任务页、独立人工标注页", "浏览器本地保存标注进度；任务数据包负责分发"],
        ["服务层", "FastAPI API、单机后台任务、签名会话认证", "评审账号是单账号演示模式，没有角色系统"],
        ["解析层", "按内容识别 DOC/XLS/PDF/Office/图片与压缩包", "压缩大小、深度、成员数均设上限"],
        ["抽取层", "规则表格解析、可选 RapidOCR、可选 Qwen/DeepSeek", "模型和 OCR 输出都是候选，不代表金标"],
        ["数据层", "SQLite 项目/包/实体/参与/中标/标的关系", "每数据集独立数据库；官方数据隔离于默认库"],
    ])

    add_heading(doc, "四 数据模型")
    add_table(doc, ["实体", "关键关系", "用途"], [
        ["organizations", "采购单位、供应商和投标主体共用组织实体；aliases 保留原名", "保守归一名称，不因相似字符串自动合并法人"],
        ["projects", "每条公告关联项目和采购单位", "项目编号、预算、公告成交总额分别保存"],
        ["packages", "项目可含多个采购包/标段", "作为投标参与和查询频次的事件单元"],
        ["bid_participations", "主体关联采购包，记录 winner/nonwinner/unknown", "保存角色、包号、来源和原文证据"],
        ["awards", "明确中标主体关联采购包", "保存主体级中标金额及证据，不以项目总额代替"],
        ["procurement_items", "标的记录关联公告与采购包", "七字段、字段价格单位、来源定位及抽取方式"],
    ])

    add_heading(doc, "五 抽取和入库流程")
    add_numbered(doc, [
        "读取上传文件并依据内容特征识别真实格式；逐层展开归档，拒绝超过安全预算的成员。",
        "读取 HTML、Office、PDF、图片文本；符合配置时用本地 OCR 处理扫描内容。失败和不可信结果保留警告。",
        "规则解析采购单位、项目头字段、标的表格、投标人/资格审查表；如启用模型，再抽取非表格内容并与规则行对齐。",
        "规范化金额和包号，去除有证据支持的重复行；对歧义和冲突保留候选并警告，不用宽松别名或排名猜测结果。",
        "以每个公告为检查点写入独立数据集；任务支持暂停、续跑、重试与导出逐条报告。",
    ])

    add_heading(doc, "六 访问控制与运行")
    add_para(doc, "本地开发默认不启用账号。演示部署可开启 AUTH_ENABLED，以强密码和随机签名密钥创建单一评审账号；后端使用 HttpOnly、SameSite=Strict 的 HMAC 签名 Cookie，过期时间默认 8 小时。配置缺失时受保护 API 返回 503。账号配置和原始数据均不进入 Git。")
    add_para(doc, "后端环境使用 Python 3.11–3.13；前端使用 Node.js/npm。安装步骤见 README，评审账号配置、登录和演示边界见 docs/REVIEWER_GUIDE.md。")

    add_heading(doc, "七 测试与实测结果")
    add_table(doc, ["验证项", "当前证据", "解释"], [
        ["后端回归", "全量测试 183 passed、1 skipped；Ruff 通过", "覆盖解析、后台任务、评测、登录与离线样例逻辑"],
        ["前端回归", "11 passed；构建通过", "构建有 ECharts 大 chunk 提示，不影响构建成功"],
        ["官方数据全量接入", "1038/1038 入库，0 最终失败；6,814 条标的候选", "rules + RapidOCR，调用模型 0 次；参与主体/中标记录为 0"],
        ["登录/离线演示", "本机登录、查询、标注页 HTTP 200；3 条虚构公告及 5 个组织", "不等于第二台评审设备网络验收"],
    ])

    add_heading(doc, "八 限制与提交前确认")
    add_bullets(doc, [
        "没有人工核验 Gold，因此不报告准确率、精确率或召回率。OCR 2261 条提示不代表识别正确。",
        "当前官方数据规则全量没有投标参与和中标记录；P1-1 新加的结构化评审表解析需在 Gold 集上验证误报和漏报。",
        "87 个无效下载附件、少量不支持格式、PDF 流错误和 1 个损坏成员仍需人工/数据方处理。",
        "Neo4j 不是当前在线 API 的关系查询后端；不可宣称线上场景使用 Neo4j 图检索增强。",
    ])
    path = OUTPUT / "system-design-draft.docx"
    doc.save(path)
    return path


def build_process_report() -> Path:
    doc = Document()
    style_document(doc, "技术路线与数据质量分析报告")
    add_cover(doc, "技术路线与数据质量分析报告", "解析方案、实测边界与评测计划")

    add_heading(doc, "一 结论与使用范围")
    add_para(doc, "本报告区分工程接入能力和抽取质量。官方全量任务证明 1038 条公告能够通过后台管线进入隔离数据库，并获得规则候选及 OCR 提示；由于任务使用 rules + RapidOCR、没有模型调用，也没有人工 Gold，它不能证明字段准确或场景查询与官方答案一致。")

    add_heading(doc, "二 技术路线对比")
    add_table(doc, ["方案", "适用内容", "速度/成本", "主要风险"], [
        ["规则解析", "明确 HTML/Office 表格、采购单位与标准标签", "本地执行，不消耗模型额度；稳定且可追溯", "正文自由叙述、扫描件和变化较大的版式覆盖有限"],
        ["Qwen/DeepSeek 模型", "非表格内容、上下文归并、结构化字段抽取", "受服务时延、并发与 token 影响；每个文档需请求", "幻觉、截断、端点参数差异及费用，需要证据校验"],
        ["Hybrid", "规则行与模型补充结果合并", "兼顾确定性与语义覆盖，仍会产生模型调用", "跨附件对齐与去重可能冲突；须保留证据并人工评测"],
        ["RapidOCR + 表格映射", "扫描图像、扫描 PDF 的文本入口", "本机推理，不消耗云模型额度", "识别文字不等于正确还原表格结构，必须抽样核验"],
    ])
    add_para(doc, "当前建议的操作次序是先用规则模式建立可复现基线，再以人工 Gold 对比规则、模型和混合方案；留出集必须与调优样本隔离。全量部署时先核对源文件哈希与解析警告，遇到错误附件不得让模型臆造缺失字段。")

    add_heading(doc, "三 数据接入与工程基线")
    add_table(doc, ["指标", "实测值", "口径"], [
        ["公告数", "1,038 / 1,038 完成；0 最终失败", "独立 SQLite 数据集，模型调用 0 次"],
        ["入口文件完整性", "2,008 个文件 SHA-256 复核；0 不匹配", "源文件未修改"],
        ["附件", "970 个外层 ZIP，合计约 3.57 GB；RAR/7z 解析出 239/13 个叶文件", "递归解析共 14,174 个叶文件，含 HTML"],
        ["抽取候选", "6,814 条标的候选；采购参与方 0、中标记录 0", "无 Gold，不代表正确率"],
        ["处理时间", "3 个工作进程，检查点估算约 66 分 45 秒；每公告中位数 1.127 秒，P95 41.249 秒", "rules + OCR 本机工程基线，不代表 model/hybrid 模式"],
        ["OCR 提示", "2,261 条识别/复核提示覆盖 359 条公告", "提示数量不是正确识别率"],
    ])

    add_heading(doc, "四 数据质量风险")
    add_table(doc, ["发现", "范围", "处理状态"], [
        ["下载错误伪附件", "87 个成员，涉及 61 条公告", "已识别并隔离；缺失内容未恢复，需数据方提供有效材料"],
        ["OCR 与可信文本不足", "2261 条 OCR 提示；另 21 条无可信文本提示覆盖 16 条公告", "应按分层样本回看原图和抽取结果"],
        ["不支持格式", "7 条格式警告，覆盖 3 条公告", "部分 WPS/专有格式需评估转换路径"],
        ["PDF 流错误", "7 条错误记录，覆盖 6 条公告", "需要人工核对文档能否重开或重新取得"],
        ["压缩成员读取错误", "1 个成员", "源内容未恢复，不能计作覆盖成功"],
        ["参与/中标关系", "全量规则候选库中均为 0", "现已加入结构化投标表规则解析；需人工验证再用于聚合"],
    ])
    add_para(doc, "这些数据描述完整性、运行告警和候选产量，不回答 TP/FP/FN。只有依据原文独立完成人工标注后，才可以按 docs/EVALUATION.md 的本地口径计算指标，并同时说明该口径与赛事官方定义的差异。")

    add_heading(doc, "五 人工验证计划")
    add_numbered(doc, [
        "使用已生成的 24 条分层标注任务包；两位标注员先共同完成 6 条试标，比较分歧并书面统一字段规则。",
        "试标分歧解决后，各自完成互不重叠的 9 条正式任务；Gold 逐条对照 HTML 与附件，自动预测只作为参考候选。",
        "统计标注一致性和遗漏类型；保留留出样本，在调优完成前不查看其模型指标。",
        "在固定软件版本、模型名、采样参数和 endpoint 下，对 rules、model、hybrid 使用完全相同的样本运行并导出报告。",
        "分别报告七字段、采购单位、投标主体、中标方/金额的精确率、召回率、F1 及约定的总体指标；按赛事正式公式复算后再填最终提交件。",
    ])

    add_heading(doc, "六 局限与后续工作")
    add_bullets(doc, [
        "官方附件有真实缺失与专有格式；无效材料不能靠 OCR 或模型恢复。",
        "没有 Gold 时，不可以把开发集、合成夹具或模型自评分包装成赛事准确率。",
        "P1-1 结构化评审表抽取要重点检查采购单位/代理机构误识别、评审分数误作金额、排名误作中标状态及包号错配。",
        "最终报告需要主办方提供指标定义、官方查询样例和评审文档模板后再冻结版本。",
    ])
    path = OUTPUT / "process-and-quality-report-draft.docx"
    doc.save(path)
    return path


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = [build_system_design(), build_process_report()]
    for result in results:
        print(result)


if __name__ == "__main__":
    main()
